#!/usr/bin/env python3

import csv
from pathlib import Path
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

PROJECT = Path("/projects/bentosprg6/linwx/research_project")

RANKED = PROJECT / "data/task5/task5_affinity_ranked_final.csv"
COVERAGE = PROJECT / "data/task5/task5_full_coverage.csv"
TARGETS = PROJECT / "data/task5/targets/target_validated.csv"

OUTPUT = PROJECT / "data/task5/TASK5_Boltz2_Aptamer_Affinity_FINAL.xlsx"


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


ranked = read_csv(RANKED)
coverage = read_csv(COVERAGE)
targets = read_csv(TARGETS)

attempted = [
    r for r in coverage
    if r["task5_attempted"] == "YES"
]

success = [
    r for r in coverage
    if r["task5_result_status"] == "SUCCESS"
]

oom = [
    r for r in coverage
    if r["task5_result_status"] == "GPU_OOM"
]

not_run = [
    r for r in coverage
    if r["task5_result_status"] == "NOT_RUN"
]

assert len(ranked) == 420
assert len(coverage) == 1495
assert len(attempted) == 425
assert len(success) == 420
assert len(oom) == 5
assert len(not_run) == 1070


# ============================================================
# Workbook
# ============================================================

wb = Workbook()

ws_summary = wb.active
ws_summary.title = "Summary"

ws_ranked = wb.create_sheet("Ranked_Affinity")
ws_attempts = wb.create_sheet("All_Attempts")
ws_coverage = wb.create_sheet("Full_Coverage")
ws_targets = wb.create_sheet("Target_Resolution")


# ============================================================
# Styles
# ============================================================

dark_fill = PatternFill("solid", fgColor="1F4E78")
blue_fill = PatternFill("solid", fgColor="D9EAF7")
green_fill = PatternFill("solid", fgColor="E2F0D9")
yellow_fill = PatternFill("solid", fgColor="FFF2CC")
red_fill = PatternFill("solid", fgColor="FCE4D6")
gray_fill = PatternFill("solid", fgColor="E7E6E6")

white_bold = Font(color="FFFFFF", bold=True)
bold = Font(bold=True)

thin_gray = Side(style="thin", color="D9E1F2")
border = Border(
    left=thin_gray,
    right=thin_gray,
    top=thin_gray,
    bottom=thin_gray
)


def style_header(ws):
    for cell in ws[1]:
        cell.fill = dark_fill
        cell.font = white_bold
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = border

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def autosize(ws, max_width=45):
    widths = {}

    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue

            length = len(str(cell.value))
            widths[cell.column] = max(
                widths.get(cell.column, 0),
                min(length + 2, max_width)
            )

    for col_idx, width in widths.items():
        ws.column_dimensions[
            get_column_letter(col_idx)
        ].width = max(10, width)


def write_dict_rows(ws, rows):
    if not rows:
        return

    fields = list(rows[0].keys())

    ws.append(fields)

    for r in rows:
        ws.append([
            r.get(field, "")
            for field in fields
        ])

    style_header(ws)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=False
            )
            cell.border = border

    autosize(ws)


def add_table(ws, name):
    if ws.max_row < 2 or ws.max_column < 1:
        return

    ref = (
        f"A1:"
        f"{get_column_letter(ws.max_column)}"
        f"{ws.max_row}"
    )

    tab = Table(
        displayName=name,
        ref=ref
    )

    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    tab.tableStyleInfo = style
    ws.add_table(tab)


# ============================================================
# Summary
# ============================================================

ws = ws_summary

ws["A1"] = "Task 5 — Boltz-2 Aptamer Affinity Prediction"
ws["A1"].font = Font(
    bold=True,
    size=16,
    color="FFFFFF"
)
ws["A1"].fill = dark_fill
ws.merge_cells("A1:D1")

ws["A3"] = "Dataset coverage"
ws["A3"].font = bold
ws["A3"].fill = blue_fill

summary_rows = [
    ("Total UTexas model records", 1495),
    ("Attempted with Boltz-2", 425),
    ("Successful affinity predictions", 420),
    ("GPU OOM", 5),
    ("Not run", 1070),
]

row = 4

for label, value in summary_rows:
    ws.cell(row=row, column=1, value=label)
    ws.cell(row=row, column=2, value=value)
    row += 1


ws["A10"] = "Successful predictions by workflow"
ws["A10"].font = bold
ws["A10"].fill = blue_fill

group_counts = Counter(
    r["result_group"]
    for r in ranked
)

ws["A11"] = "Original direct-ready"
ws["B11"] = group_counts["ORIGINAL_DIRECT_READY"]

ws["A12"] = "Recovery — unmodified canonical backbone"
ws["B12"] = group_counts[
    "RECOVERY_UNMODIFIED_BACKBONE"
]


ws["A14"] = "Full coverage disposition"
ws["A14"].font = bold
ws["A14"].fill = blue_fill

coverage_counts = Counter(
    r["task5_coverage_group"]
    for r in coverage
)

r = 15

for status, count in coverage_counts.most_common():
    ws.cell(row=r, column=1, value=status)
    ws.cell(row=r, column=2, value=count)
    r += 1


ws["A24"] = "Prediction range"
ws["A24"].font = bold
ws["A24"].fill = blue_fill

lowest = ranked[0]
highest = ranked[-1]

ws["A25"] = "Lowest affinity_pred_value"
ws["B25"] = float(lowest["affinity_pred_value"])
ws["C25"] = lowest["model_id"]
ws["D25"] = lowest["target_normalized"]

ws["A26"] = "Highest affinity_pred_value"
ws["B26"] = float(highest["affinity_pred_value"])
ws["C26"] = highest["model_id"]
ws["D26"] = highest["target_normalized"]


ws["A28"] = "Important interpretation notes"
ws["A28"].font = bold
ws["A28"].fill = yellow_fill

notes = [
    (
        "Boltz-2 was adapted so DNA/RNA polymer chains could be "
        "specified as the affinity binder."
    ),
    (
        "83 additional records with canonical DNA/RNA backbones "
        "were modeled as an unmodified-backbone approximation."
    ),
    (
        "Chemical modifications present in the experimental "
        "aptamers are not explicitly represented in those recovery runs."
    ),
    (
        "Affinity outputs should not be interpreted as experimentally "
        "validated or quantitatively reliable aptamer binding affinities."
    ),
    (
        "Five original oversized protein complexes failed structure "
        "prediction because of GPU-memory limits and produced no "
        "affinity value."
    ),
    (
        "No model retraining was performed."
    ),
]

for i, note in enumerate(notes, 29):
    ws.cell(row=i, column=1, value="• " + note)
    ws.merge_cells(
        start_row=i,
        start_column=1,
        end_row=i,
        end_column=4
    )
    ws.cell(row=i, column=1).alignment = Alignment(
        wrap_text=True,
        vertical="top"
    )


for row_cells in ws.iter_rows(
    min_row=3,
    max_row=ws.max_row,
    min_col=1,
    max_col=4
):
    for cell in row_cells:
        cell.border = border
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )

ws.column_dimensions["A"].width = 48
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 28
ws.column_dimensions["D"].width = 55


# ============================================================
# Ranked Affinity — 420 successes
# ============================================================

write_dict_rows(ws_ranked, ranked)
add_table(ws_ranked, "RankedAffinityTable")

# Affinity column number format
headers = {
    cell.value: cell.column
    for cell in ws_ranked[1]
}

if "affinity_pred_value" in headers:
    col = get_column_letter(
        headers["affinity_pred_value"]
    )

    for cell in ws_ranked[col][1:]:
        try:
            cell.value = float(cell.value)
            cell.number_format = "0.000000"
        except Exception:
            pass


# ============================================================
# All Attempts — 425
# ============================================================

attempt_rows = []

for r in attempted:
    attempt_rows.append(r)

write_dict_rows(ws_attempts, attempt_rows)
add_table(ws_attempts, "AllAttemptsTable")

status_col = None

for cell in ws_attempts[1]:
    if cell.value == "task5_result_status":
        status_col = cell.column
        break

if status_col:
    for row_idx in range(2, ws_attempts.max_row + 1):
        cell = ws_attempts.cell(
            row=row_idx,
            column=status_col
        )

        if cell.value == "SUCCESS":
            cell.fill = green_fill

        elif cell.value == "GPU_OOM":
            cell.fill = red_fill


# ============================================================
# Full Coverage — 1495
# ============================================================

write_dict_rows(ws_coverage, coverage)
add_table(ws_coverage, "FullCoverageTable")

coverage_col = None

for cell in ws_coverage[1]:
    if cell.value == "task5_coverage_group":
        coverage_col = cell.column
        break

if coverage_col:
    for row_idx in range(2, ws_coverage.max_row + 1):
        cell = ws_coverage.cell(
            row=row_idx,
            column=coverage_col
        )

        value = str(cell.value or "")

        if value.startswith("SUCCESS"):
            cell.fill = green_fill

        elif value.startswith("GPU_OOM"):
            cell.fill = red_fill

        elif "REVIEW" in value:
            cell.fill = yellow_fill

        else:
            cell.fill = gray_fill


# ============================================================
# Target Resolution
# ============================================================

write_dict_rows(ws_targets, targets)
add_table(ws_targets, "TargetResolutionTable")


# ============================================================
# Workbook metadata / final save
# ============================================================

wb.properties.title = \
    "Task 5 Boltz-2 Aptamer Affinity Results"

wb.properties.subject = \
    "UTexas aptamer affinity prediction workflow"

wb.properties.creator = \
    "Task 5 research workflow"

wb.save(OUTPUT)

print("========== FINAL TASK 5 EXCEL CREATED ==========")
print("Workbook:", OUTPUT)

print()
print("Sheets:")
for name in wb.sheetnames:
    ws = wb[name]
    print(
        f"{name:20s} "
        f"rows={ws.max_row:5d} "
        f"cols={ws.max_column:3d}"
    )

print()
print("Ranked successful predictions :", len(ranked))
print("All attempted complexes       :", len(attempted))
print("Full UTexas coverage          :", len(coverage))
print("Target-resolution rows        :", len(targets))

print()
print("========== EXCEL CREATION PASSED ==========")
